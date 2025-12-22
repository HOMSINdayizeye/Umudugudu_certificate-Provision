from urllib import request
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login as auth_login, logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from .forms import SignUpForm, CertificateRequestForm, StolenLaptopCertificateForm
from .models import EligiblePerson, CertificateRequest, CustomUser, LocationImport, Location, StolenLaptopCertificate
from django.http import HttpResponseForbidden, HttpResponseBadRequest, JsonResponse
from django.conf import settings
from django.http import HttpResponse
from docx import Document
import io
import datetime 
import os



def index(request):
    return render(request, 'certificates/index.html')


def register(request):
    if request.method == 'POST':
        form = SignUpForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.email = form.cleaned_data['email']
            user.first_name = form.cleaned_data['first_name']
            user.last_name = form.cleaned_data['last_name']
            user.save()
            messages.success(request, 'Account created successfully!')
            return redirect('login')
    else:
        form = SignUpForm()
    
    return render(request, 'certificates/register.html', {'form': form})


@login_required
def choose_certificate(request):
    return render(request, 'certificates/choose_certificate.html')


@login_required
def submit_request(request):
    if request.method == 'POST':
        form = CertificateRequestForm(request.POST)
        
        if form.is_valid():
            # Create the certificate request
            cert_request = form.save(commit=False)
            cert_request.user = request.user
            cert_request.status = 'pending'
            
            cert_type = form.cleaned_data.get('cert_type')
            
            # Handle location code for conduct certificates
            if cert_type == 'conduct':
                location_code = request.POST.get('location_code')
                if location_code:
                    cert_request.location_code = location_code
            
            cert_request.save()
            
            # Handle stolen laptop/computer certificate data
            if cert_type in ['stolen_laptop', 'stolen_computer']:
                StolenLaptopCertificate.objects.create(
                    certificate_request=cert_request,
                    registration_number=form.cleaned_data.get('registration_number'),
                    national_id=form.cleaned_data.get('national_id'),
                    stolen_datetime=form.cleaned_data.get('stolen_datetime'),
                    location=form.cleaned_data.get('stolen_location'),
                    device_type=form.cleaned_data.get('device_type'),
                    brand=form.cleaned_data.get('brand'),
                    model=form.cleaned_data.get('model'),
                    serial_number=form.cleaned_data.get('serial_number'),
                    intel_core=form.cleaned_data.get('intel_core'),
                    ram=form.cleaned_data.get('ram'),
                    storage=form.cleaned_data.get('storage'),
                    witness_1_name=form.cleaned_data.get('witness_1_name'),
                    witness_1_phone=form.cleaned_data.get('witness_1_phone'),
                    witness_2_name=form.cleaned_data.get('witness_2_name', ''),
                    witness_2_phone=form.cleaned_data.get('witness_2_phone', ''),
                    witness_3_name=form.cleaned_data.get('witness_3_name', ''),
                    witness_3_phone=form.cleaned_data.get('witness_3_phone', ''),
                    reported_to_police=form.cleaned_data.get('reported_to_police', False),
                    other_description=form.cleaned_data.get('laptop_other_description', '')
                )
            
            messages.success(request, 'Certificate request submitted successfully!')
            return redirect('dashboard')
        else:
            messages.error(request, 'Please correct the errors below.')
    else:
        form = CertificateRequestForm()
    
    # FIXED: Get provinces from LocationImport model
    provinces = LocationImport.objects.filter(
        type='PROVINCE'
    ).values('location_id', 'name').order_by('location_id')
    
    context = {
        'form': form,
        'provinces': list(provinces)  # Convert QuerySet to list
    }
    return render(request, 'certificates/submit_request.html', context)


@login_required
def request_detail(request, pk):
    if request.user.is_superuser:
        req = get_object_or_404(CertificateRequest, pk=pk)
    else:
        req = get_object_or_404(CertificateRequest, pk=pk, user=request.user)
    return render(request, 'certificates/request_detail.html', {'req': req})


@user_passes_test(lambda u: u.is_superuser)
def update_request_status(request, pk, action):
    req = get_object_or_404(CertificateRequest, pk=pk)
    if action == 'approve':
        req.status = 'approved'
        req.admin_message = 'Your request has been approved.'
    elif action == 'deny':
        req.status = 'denied'
        req.admin_message = 'Your request has been denied.'
    else:
        return HttpResponseBadRequest('Invalid action.')
    req.save()
    messages.success(request, f'Request has been {action}ed successfully.')
    return redirect('dashboard')


@user_passes_test(lambda u: u.is_superuser)
def manage_eligibility(request):
    users = CustomUser.objects.all()
    if request.method == 'POST':
        eligible_ids = request.POST.getlist('eligible')
        for user in users:
            user.is_eligible = str(user.id) in eligible_ids
            user.save()
        messages.success(request, 'Eligibility updated successfully!')
        return redirect('manage_eligibility')
    return render(request, 'certificates/manage_eligibility.html', {'users': users})


def login_view(request):
    if request.method == 'POST':
        username = request.POST.get('email')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            auth_login(request, user)
            messages.success(request, 'Login successful!')
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid email or password.')
    return render(request, 'certificates/login.html')


def logout_view(request):
    logout(request)
    messages.success(request, 'You have been logged out successfully.')
    return redirect('index')


@login_required
def dashboard(request):
    """Dashboard view for users and admins"""
    if request.user.is_superuser:
        # Admin dashboard - show all requests
        requests = CertificateRequest.objects.all().order_by('-created_at')
    else:
        # User dashboard - show only their requests
        requests = CertificateRequest.objects.filter(user=request.user).order_by('-created_at')
    
    context = {
        'requests': requests
    }
    return render(request, 'certificates/dashboard.html', context)


# ==================== AJAX LOCATION FILTERING VIEWS ====================

def get_provinces(request):
    """Get all provinces"""
    provinces = LocationImport.objects.filter(
        type='PROVINCE'
    ).values('location_id', 'name').order_by('location_id')
    
    return JsonResponse({'provinces': list(provinces)})


def get_districts(request):
    """Get districts based on selected province"""
    province_id = request.GET.get('province')
    
    if not province_id:
        return JsonResponse({'districts': []})
    
    province_id_str = str(province_id)
    
    districts = LocationImport.objects.filter(
        type='DISTRICT'
    ).extra(
        where=[
            "CAST(location_id AS TEXT) LIKE %s",
            "CHAR_LENGTH(CAST(location_id AS TEXT)) = 2"
        ],
        params=[f"{province_id_str}%"]
    ).values('location_id', 'name').order_by('location_id')
    
    return JsonResponse({'districts': list(districts)})


def get_sectors(request):
    """Get sectors based on selected district"""
    district_id = request.GET.get('district')
    
    if not district_id:
        return JsonResponse({'sectors': []})
    
    district_id_str = str(district_id)
    
    sectors = LocationImport.objects.filter(
        type='SECTOR'
    ).extra(
        where=[
            "CAST(location_id AS TEXT) LIKE %s",
            "CHAR_LENGTH(CAST(location_id AS TEXT)) = 4"
        ],
        params=[f"{district_id_str}%"]
    ).values('location_id', 'name').order_by('location_id')
    
    return JsonResponse({'sectors': list(sectors)})


def get_cells(request):
    """Get cells based on selected sector"""
    sector_id = request.GET.get('sector')
    
    if not sector_id:
        return JsonResponse({'cells': []})
    
    sector_id_str = str(sector_id)
    
    cells = LocationImport.objects.filter(
        type='CELL'
    ).extra(
        where=[
            "CAST(location_id AS TEXT) LIKE %s",
            "CHAR_LENGTH(CAST(location_id AS TEXT)) = 6"
        ],
        params=[f"{sector_id_str}%"]
    ).values('location_id', 'name').order_by('location_id')
    
    return JsonResponse({'cells': list(cells)})


def get_villages(request):
    """Get villages based on selected cell"""
    cell_id = request.GET.get('cell')
    
    if not cell_id:
        return JsonResponse({'villages': []})
    
    cell_id_str = str(cell_id)
    
    # If you don't have villages in your data, this will return empty
    villages = LocationImport.objects.filter(
        type='VILLAGE'
    ).extra(
        where=[
            "CAST(location_id AS TEXT) LIKE %s",
            "CHAR_LENGTH(CAST(location_id AS TEXT)) >= 7"
        ],
        params=[f"{cell_id_str}%"]
    ).values('location_id', 'name').order_by('location_id')
    
    return JsonResponse({'villages': list(villages)})


def get_location_code(request):
    """Get final location code based on selection"""
    village_id = request.GET.get('village')
    cell_id = request.GET.get('cell')
    sector_id = request.GET.get('sector')
    district_id = request.GET.get('district')
    province_id = request.GET.get('province')
    
    location_id = village_id or cell_id or sector_id or district_id or province_id
    
    if location_id:
        try:
            location = LocationImport.objects.get(location_id=location_id)
            return JsonResponse({
                'code': location.location_id,
                'name': location.name,
                'type': location.type,
                'success': True
            })
        except LocationImport.DoesNotExist:
            return JsonResponse({
                'code': None,
                'success': False,
                'error': 'Location not found'
            })
    
    return JsonResponse({
        'code': None,
        'success': False,
        'error': 'No location selected'
    })


@login_required
def report_stolen_laptop(request):
    form = StolenLaptopCertificateForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        form.save()
        return redirect('dashboard')
    return render(request, 'send_request.html', {
        'form': form,
        'form_title': 'Report Stolen Laptop',
        'submit_label': 'Submit Report',
    })
# function to display initialsgit
def display_user_info(firstname, lastname):
    # Get initials
    initials = firstname[0].upper() + lastname[0].upper()
    
    # Welcome message
    welcome_message = f"Welcome back {firstname}"
    
    # Print results
    print(f"FIRSTNAME: {firstname}, LASTNAME: {lastname}")
    print(welcome_message)
    print(f"Initials in circle: ({initials})")
    #cerificate function from templates /////////////////
def certificate_request_view(request):
    if request.method == "POST":
        template_path = os.path.join(
            settings.BASE_DIR, "certificates", "templates_docs", "conduct_template.docx"
        )
        doc = Document(template_path)

        # Collect form data safely
        name = request.POST.get("name", "")
        nid = request.POST.get("nid", "")
        village = request.POST.get("village_name", "")
        father = request.POST.get("father_name", "")
        mother = request.POST.get("mother_name", "")
        province = request.POST.get("province_name", "")
        district = request.POST.get("district_name", "")
        sector = request.POST.get("sector_name", "")
        cell = request.POST.get("cell_name", "")
        dob = request.POST.get("dob", "")
        district_issue = request.POST.get("district_of_issue", "")
        sector_issue = request.POST.get("id_issue_sector", "")
        date = request.POST.get("date", "")
        leader = request.POST.get("leader_name", "")

        replacements = {
            "{{name}}": str(name),
            "{{nid}}": str(nid),
            "{{village_name}}": str(village),
            "{{father_name}}": str(father),
            "{{mother_name}}": str(mother),
            "{{province_name}}": str(province),
            "{{district_name}}": str(district),
            "{{sector_name}}": str(sector),
            "{{cell_name}}": str(cell),
            "{{dob}}": str(dob),
            "{{district_of_issue}}": str(district_issue),
            "{{id_issue_sector}}": str(sector_issue),
            "{{date}}": str(date),
            "{{leader_name}}": str(leader),
        }

        # Replace placeholders in runs
        for p in doc.paragraphs:
            for run in p.runs:
                for key, val in replacements.items():
                    if key in run.text:
                        run.text = run.text.replace(key, val)

        # Save to memory and return
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)

        response = HttpResponse(
            output_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        response['Content-Disposition'] = 'attachment; filename=conduct_certificate.docx'
        return response

    return render(request, "certificate_request.html")
#generating certificate////////////////////////////
@login_required
def generate_certificate(request, pk):
    # Only allow POST
    if request.method == "POST":
        cert = get_object_or_404(CertificateRequest, pk=pk)

        # Optionally mark as approved
        cert.status = "approved"
        cert.save()

        # Load template
        template_path = os.path.join(
            settings.BASE_DIR, "certificates", "templates_docs", "conduct_template.docx"
        )
        doc = Document(template_path)

        # Replace placeholders with model fields
        replacements = {
            "{{name}}": str(cert.name),
            "{{nid}}": str(cert.nid),
            "{{village_name}}": str(cert.village_name),
            # add other fields from your model here
        }

        for p in doc.paragraphs:
            for run in p.runs:
                for key, val in replacements.items():
                    if key in run.text:
                        run.text = run.text.replace(key, val)

        # Return as download
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)

        response = HttpResponse(
            output_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        response['Content-Disposition'] = 'attachment; filename=certificate.docx'
        return response

    return HttpResponse("Invalid request", status=400)