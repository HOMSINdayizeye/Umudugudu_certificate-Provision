"""Builds the official village letters (.docx) in the same layout as the Isibo office documents."""
import datetime
import io

from django.utils import timezone
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .models import LocationImport

RW_MONTHS = ['Mutarama', 'Gashyantare', 'Werurwe', 'Mata', 'Gicurasi', 'Kamena',
             'Nyakanga', 'Kanama', 'Nzeri', 'Ukwakira', 'Ugushyingo', 'Ukuboza']
EN_PROVINCE = {
    'umujyi wa kigali': 'CITY OF KIGALI',
    'amajyaruguru': 'NORTHERN PROVINCE',
    'amajyepfo': 'SOUTHERN PROVINCE',
    'iburasirazuba': 'EASTERN PROVINCE',
    'iburengerazuba': 'WESTERN PROVINCE',
}
VOWELS = 'aeiouAEIOU'


# ---------- helpers ----------

def location_chain(location_id):
    """Names of province → village for a hierarchical location id (1/2/4/6/8 digit prefixes)."""
    names = {'province': '', 'district': '', 'sector': '', 'cell': '', 'village': ''}
    if not location_id:
        return names
    s = str(location_id)
    for key, n in (('province', 1), ('district', 2), ('sector', 4), ('cell', 6), ('village', 8)):
        if len(s) >= n:
            loc = LocationImport.objects.filter(location_id=int(s[:n])).first()
            if loc:
                names[key] = loc.name
    return names


def rw_of(word, stem):
    """Kinyarwanda possessive: "w'Isibo" before a vowel, "wa Kabuye" otherwise (case follows stem)."""
    if not word:
        return stem
    if word[0] in VOWELS:
        return f"{stem}'{word}"
    return f"{stem}{'A' if stem.isupper() else 'a'} {word}"


def parse_date(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value
    if not value:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m'):
        try:
            return datetime.datetime.strptime(str(value), fmt).date()
        except ValueError:
            continue
    return None


def rw_date(value):
    d = parse_date(value)
    if not d:
        return str(value or '')
    return f'{d.day:02d} {RW_MONTHS[d.month - 1]} {d.year}'


def rw_month_year(value):
    d = parse_date(value)
    if not d:
        return str(value or '')
    return f'{RW_MONTHS[d.month - 1]} {d.year}'


def en_date(value):
    d = parse_date(value)
    if not d:
        return str(value or '')
    return d.strftime('%B %d, %Y')


def en_long_date(value):
    d = parse_date(value)
    if not d:
        return str(value or '')
    return d.strftime('%A, %d %B %Y')


def pronouns(gender):
    g = (gender or '').lower()
    if g == 'male':
        return {'he': 'he', 'his': 'his', 'him': 'him', 'title': 'Mr.'}
    if g == 'female':
        return {'he': 'she', 'his': 'her', 'him': 'her', 'title': 'Ms.'}
    return {'he': 'they', 'his': 'their', 'him': 'them', 'title': ''}


def _new_document():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    for section in doc.sections:
        # A4 portrait, stated explicitly so every viewer agrees
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    return doc


def docx_to_pdf(docx_bytes):
    """Render one of our letters to PDF: paragraphs with bold/italic runs, bullets and the letterhead table."""
    from xml.sax.saxutils import escape
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    source = Document(io.BytesIO(docx_bytes))
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title='Letter')
    base = ParagraphStyle('base', fontName='Times-Roman', fontSize=12, leading=16)
    align = {WD_ALIGN_PARAGRAPH.CENTER: TA_CENTER, WD_ALIGN_PARAGRAPH.RIGHT: TA_RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY: TA_JUSTIFY}

    def markup(p):
        parts = []
        for r in p.runs:
            t = escape(r.text)
            if r.bold:
                t = f'<b>{t}</b>'
            if r.italic:
                t = f'<i>{t}</i>'
            if r.underline:
                t = f'<u>{t}</u>'
            parts.append(t)
        return ''.join(parts) or escape(p.text)

    def para_style(p, space_after=None):
        sa = p.paragraph_format.space_after
        return ParagraphStyle('p', parent=base, alignment=align.get(p.alignment, TA_LEFT),
                              spaceAfter=space_after if space_after is not None else (sa.pt if sa is not None else 8))

    story = []
    for child in source.element.body.iterchildren():
        if child.tag.endswith('}tbl'):
            table = DocxTable(child, source)
            data = [[[Paragraph(markup(p), para_style(p, 0)) for p in cell.paragraphs] for cell in row.cells]
                    for row in table.rows]
            ncols = len(table.columns)
            widths = [pdf.width * 0.62, pdf.width * 0.38] if ncols == 2 else [pdf.width / ncols] * ncols
            t = Table(data, colWidths=widths)
            t.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                   ('LEFTPADDING', (0, 0), (-1, -1), 0), ('RIGHTPADDING', (0, 0), (-1, -1), 0)]))
            story.append(t)
        elif child.tag.endswith('}p'):
            p = DocxParagraph(child, source)
            if not p.text.strip():
                story.append(Spacer(1, 12))
                continue
            style = para_style(p)
            if p.style is not None and 'List' in p.style.name:
                style.leftIndent = 18
                style.bulletIndent = 6
                story.append(Paragraph(markup(p), style, bulletText='•'))
            else:
                story.append(Paragraph(markup(p), style))
    pdf.build(story)
    return buf.getvalue()


def _letterhead(doc, chain, date_text, lang):
    if lang == 'rw':
        province = chain['province']
        first = 'UMUJYI WA KIGALI' if 'kigali' in province.lower() else rw_of(province.upper(), "INTARA Y")
        lines = [
            first,
            f"AKARERE KA {chain['district'].upper()}",
            f"UMURENGE WA {chain['sector'].upper()}",
            f"AKAGARI KA {chain['cell'].upper()}",
            rw_of(chain['village'].upper(), 'UMUDUGUDU W'),
        ]
    else:
        lines = [
            EN_PROVINCE.get(chain['province'].lower(), f"{chain['province'].upper()} PROVINCE"),
            f"{chain['district'].upper()} DISTRICT",
            f"{chain['sector'].upper()} SECTOR",
            f"{chain['cell'].upper()} CELL",
            f"{chain['village'].upper()} VILLAGE",
        ]
    table = doc.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    for i, line in enumerate(lines):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.bold = True
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run(date_text)
    doc.add_paragraph()


def _title(doc, text, underline=False, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.underline = underline
    p.paragraph_format.space_after = Pt(10)
    return p


def _para(doc, parts, align_justify=True, space_after=8):
    """Add a paragraph from a string or a list of (text, bold) tuples."""
    p = doc.add_paragraph()
    if align_justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if isinstance(parts, str):
        parts = [(parts, False)]
    for text, bold in parts:
        run = p.add_run(text)
        run.bold = bold
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)


def _signature(doc, title, name, phone=''):
    doc.add_paragraph()
    p = doc.add_paragraph(title)
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.add_run(name).bold = True
    p.paragraph_format.space_after = Pt(0)
    if phone:
        doc.add_paragraph(f'Tel: {phone}')


def _blank_lines(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _g(details, key, default=''):
    value = details.get(key, default)
    return default if value is None else str(value).strip()


def birthplace_rw(d):
    """"yavukiye mu Ntara y'…, Akarere ka …, Umurenge wa …, Akagari ka …, Umudugudu wa …" from free-text fields."""
    parts = []
    country = _g(d, 'birth_country')
    province = _g(d, 'birth_province')
    if country and 'rwanda' not in country.lower():
        parts.append(f"mu gihugu cya {country}")
    if province:
        if 'kigali' in province.lower():
            parts.append('mu Mujyi wa Kigali')
        else:
            parts.append(rw_of(province, "mu Ntara y"))
    if _g(d, 'birth_district'):
        parts.append(f"mu Karere ka {_g(d, 'birth_district')}")
    if _g(d, 'birth_sector'):
        parts.append(f"Umurenge wa {_g(d, 'birth_sector')}")
    if _g(d, 'birth_cell'):
        parts.append(f"Akagari ka {_g(d, 'birth_cell')}")
    if _g(d, 'birth_village'):
        parts.append(rw_of(_g(d, 'birth_village'), 'Umudugudu w'))
    return ', '.join(parts)


# ---------- request letters ----------

def build_conduct(req, d, chain, leader, today):
    doc = _new_document()
    _letterhead(doc, chain, rw_date(today), 'rw')
    village = chain['village']
    _title(doc, "IMPAMVU: Imico n'imyifatire")

    name = _g(d, 'full_name')
    school = _g(d, 'school')
    department = _g(d, 'department')
    year = _g(d, 'year_of_study')
    occupation = _g(d, 'occupation')

    first = [(f"Ubuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'bw')} tubandikiye tubamenyesha ko uyu witwa ", False),
             (name, True),
             (f" tumuzi abarizwa mu mudugudu {rw_of(village, 'w')}", False)]
    if _g(d, 'occupation_type', 'student') == 'student':
        if school:
            first.append((f", ubarizwa mu ishuri rya {school}", False))
        if department:
            first.append((f", akaba ari umunyeshuri mu bijyanye na {department}", False))
        if year:
            first.append((f", umwaka wa {year}", False))
    elif occupation:
        first.append((f", akaba ari {occupation}", False))
    first.append((". Turabasaba ko mwamufasha kubona icyangombwa cy'imico n'imyifatire gitangwa n'akagari, "
                  "kuko tumuzi asanzwe ari indakemwa mu mico n'imyifatire.", False))
    _para(doc, first)

    second = [("Ni mwene ", False), (_g(d, 'father_name'), True), (" na ", False), (_g(d, 'mother_name'), True)]
    place = birthplace_rw(d)
    if place:
        second.append((f", yavukiye {place}", False))
    second.append((f". Yavutse kuwa {rw_date(_g(d, 'dob'))}, afite nomero y'indangamuntu ", False))
    second.append((_g(d, 'national_id'), True))
    issue = []
    if _g(d, 'id_issue_district'):
        issue.append(f"mu Karere ka {_g(d, 'id_issue_district')}")
    if _g(d, 'id_issue_sector'):
        issue.append(f"Umurenge wa {_g(d, 'id_issue_sector')}")
    if issue:
        second.append((f", yatangiwe {', '.join(issue)}", False))
    second.append(('.', False))
    _para(doc, second)

    if _g(d, 'purpose'):
        _para(doc, f"Icyangombwa agisabira: {_g(d, 'purpose')}.")

    _para(doc, 'Mumwakire murakoze!', align_justify=False)
    _signature(doc, f"Umuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'w')}", leader['name'], leader['phone'])
    return _save(doc)


def build_residence(req, d, chain, leader, today):
    doc = _new_document()
    _letterhead(doc, chain, f'Kuwa {rw_date(today)}', 'rw')
    village = chain['village']
    _title(doc, 'ICYANGOMBWA KIGARAGAZA KO UMUTURAGE AZWI', center=True)

    parts = [(f"Ubuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'bw')} tubandikiye tubamenyesha ko uyu witwa ", False),
             (_g(d, 'full_name'), True),
             (' ufite Indangamuntu nomero: ', False), (_g(d, 'national_id'), True)]
    issue = ' / '.join(x for x in (_g(d, 'id_issue_district'), _g(d, 'id_issue_sector')) if x)
    if issue:
        parts.append((f', yatangiwe {issue}', False))
    parts.append((f", yavutse kuwa {rw_date(_g(d, 'dob'))}, mwene ", False))
    parts.append((_g(d, 'father_name'), True))
    parts.append((' na ', False))
    parts.append((_g(d, 'mother_name'), True))
    place = birthplace_rw(d)
    if place:
        parts.append((f', yavukiye {place}', False))
    parts.append(('.', False))
    _para(doc, parts)

    since = rw_month_year(_g(d, 'resident_since'))
    _para(doc, f"Ni umuturage mu mudugudu {rw_of(village, 'w')} kuva {since}, kandi azwi nk'umuturage utuye aha "
               f"mu mudugudu {rw_of(village, 'w')}.")
    _para(doc, "Icyangombwa gitanzwe n'ubuyobozi bw'umudugudu gishobora kwifashishwa mu kugaragaza ko umuturage "
               "atuye ahavuzwe haruguru.")

    valid_until = today + datetime.timedelta(days=30)
    _para(doc, f'Gitanzwe kuwa: {rw_date(today)}', align_justify=False, space_after=0)
    _para(doc, f'Gishobora gukoreshwa bitarenze: {rw_date(valid_until)}', align_justify=False)
    _signature(doc, f"Umuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'w')}", leader['name'], leader['phone'])
    return _save(doc)


def build_stolen_device(req, d, chain, leader, today):
    doc = _new_document()
    _letterhead(doc, chain, en_date(today), 'en')
    village = chain['village']
    pr = pronouns(_g(d, 'gender'))
    device = _g(d, 'device_type', 'Laptop')
    _title(doc, f'Report of Stolen {device}', center=True)

    intro = [(f'I am writing to formally report the theft of a {device.lower()} belonging to ', False),
             (_g(d, 'full_name'), True)]
    if _g(d, 'nationality'):
        intro.append((f", nationality: {_g(d, 'nationality')}", False))
    if _g(d, 'registration_number'):
        institution = _g(d, 'institution') or 'the University of Rwanda'
        intro.append((f", with Registration number {_g(d, 'registration_number')} provided by {institution}", False))
    id_label = 'passport number' if _g(d, 'id_type') == 'passport' else 'national identification number'
    intro.append((f', a citizen residing at {village} village, whose {id_label} is ', False))
    intro.append((_g(d, 'id_number'), True))
    intro.append(('.', False))
    _para(doc, intro)

    when = en_date(_g(d, 'incident_date'))
    if _g(d, 'incident_time'):
        when += f" at approximately {_g(d, 'incident_time')}"
    spec = [x for x in (_g(d, 'brand'), _g(d, 'model')) if x]
    extras = [x for x in (_g(d, 'processor'), f"{_g(d, 'ram')} RAM" if _g(d, 'ram') else '',
                          _g(d, 'storage'), f"colour {_g(d, 'color')}" if _g(d, 'color') else '') if x]
    description = ' '.join(spec)
    if extras:
        description += (', ' if description else '') + ', '.join(extras)
    body = (f"According to the complainant, the incident occurred on {when}, at {_g(d, 'incident_location')}. "
            f"The {device.lower()} was last seen in {pr['his']} possession at that location, and upon returning, "
            f"{pr['he']} discovered it missing. A preliminary check of the surroundings was conducted. ")
    body_parts = [(body, False), (f'The stolen device is described as: {description}', False)]
    if _g(d, 'serial_number'):
        body_parts.append((', Serial Number ', False))
        body_parts.append((_g(d, 'serial_number'), True))
    body_parts.append(('.', False))
    _para(doc, body_parts)

    if _g(d, 'other_description'):
        _para(doc, _g(d, 'other_description'))

    witnesses = [w for w in (d.get('witnesses') or []) if isinstance(w, dict) and str(w.get('name', '')).strip()]
    if witnesses:
        _para(doc, 'Upon discovering the missing device, the complainant immediately informed the following '
                   'persons, who are available to provide a statement regarding the incident:', space_after=2)
        _bullets(doc, [f"{w['name'].strip()}" + (f" ({str(w.get('phone', '')).strip()})" if str(w.get('phone', '')).strip() else '')
                       for w in witnesses])

    reported = _g(d, 'reported_to')
    if d.get('reported_to_police'):
        reported = (reported + ', ' if reported else '') + 'Rwanda National Police'
    if reported:
        _para(doc, f'The incident was also reported to: {reported}.')

    _para(doc, 'This document is provided upon request of the complainant as our citizen. The complainant has '
               'also been advised to notify their service provider and monitor any attempts to access or resell '
               'the device. This letter is issued upon the request of the complainant for the purposes of official '
               'follow-up, or any other lawful use.')
    _signature(doc, f'{village} Village Leader', leader['name'], leader['phone'])
    return _save(doc)


def build_community(req, d, chain, leader, today):
    doc = _new_document()
    _letterhead(doc, chain, en_date(today), 'en')
    pr = pronouns(_g(d, 'gender'))
    name = _g(d, 'full_name')
    surname = name.split()[0] if name else ''
    title_name = f"{pr['title']} {surname}".strip() if pr['title'] else name

    _para(doc, 'To Whom It May Concern,', align_justify=False)
    _title(doc, 'RE: COMMUNITY ENGAGEMENT REFERRAL LETTER')
    _para(doc, [('This is to certify that ', False), (f"{pr['title']} {name}".strip(), True),
                (', holder of National ID number ', False), (_g(d, 'national_id'), True),
                (f", is a resident of {chain['village']} village, {chain['cell']} cell, {chain['sector']} sector, "
                 f"{chain['district']} district.", False)])
    activities = _g(d, 'activities') or 'Umuganda, youth programs, public talks and local development initiatives'
    _para(doc, f"{title_name} has been actively engaged in community activities, including: {activities}. "
               f"{pr['his'].capitalize()} participation has been consistent, responsible, and in good standing "
               f"with local leadership.")
    _para(doc, f"As the Village Leader, I confirm that {title_name} is known for {pr['his']} positive contribution "
               f"to the community and is eligible for any opportunity that requires proof of community engagement.")
    if _g(d, 'purpose'):
        _para(doc, f"This letter is issued for the purpose of: {_g(d, 'purpose')}.")
    _para(doc, 'Should you require any further information, please do not hesitate to contact our office.')
    _para(doc, 'Sincerely,', align_justify=False)
    _signature(doc, 'Village Leader', leader['name'], leader['phone'])
    return _save(doc)


def build_other(req, d, chain, leader, today):
    doc = _new_document()
    _letterhead(doc, chain, en_date(today), 'en')
    _para(doc, 'To Whom It May Concern,', align_justify=False)
    _title(doc, f"RE: {_g(d, 'subject') or 'ATTESTATION'}".upper())
    applicant = _g(d, 'full_name') or req.user.display_name
    _para(doc, [('This is to certify that ', False), (applicant, True),
                (f", is a resident of {chain['village']} village, {chain['cell']} cell, {chain['sector']} sector, "
                 f"{chain['district']} district, and is known to the village leadership.", False)])
    if req.other_description:
        _para(doc, req.other_description)
    _para(doc, 'Should you require any further information, please do not hesitate to contact our office.')
    _para(doc, 'Sincerely,', align_justify=False)
    _signature(doc, 'Village Leader', leader['name'], leader['phone'])
    return _save(doc)


BUILDERS = {
    'conduct': build_conduct,
    'residence': build_residence,
    'stolen_computer': build_stolen_device,
    'community': build_community,
    'other': build_other,
}

FILE_SLUGS = {
    'conduct': 'Certificate_of_Conduct',
    'residence': 'Residence_Recognition',
    'stolen_computer': 'Stolen_Device_Report',
    'community': 'Community_Engagement_Referral',
    'other': 'Attestation',
}


def leader_info(user):
    if not user:
        return {'name': '', 'phone': ''}
    return {'name': user.display_name, 'phone': user.phone or ''}


def request_location_id(req):
    return (req.village or req.location_code or req.cell or req.sector or req.district
            or req.province or req.user.village)


def render_request_document(req, leader=None):
    """Return (docx bytes, filename) for an approved request."""
    details = req.details if isinstance(req.details, dict) else {}
    chain = location_chain(request_location_id(req))
    signer = leader or req.approved_by
    today = timezone.localtime(req.approved_at).date() if req.approved_at else datetime.date.today()
    builder = BUILDERS.get(req.cert_type, build_other)
    data = builder(req, details, chain, leader_info(signer), today)
    applicant = (details.get('full_name') or req.user.display_name or 'applicant').replace(' ', '_')
    filename = f"{FILE_SLUGS.get(req.cert_type, 'Document')}_{applicant}_{req.pk}.docx"
    return data, filename


# ---------- announcements ----------

def announcement_paragraphs(a, chain, leader_name):
    """Paragraphs for an announcement as (style, text) tuples; shared by preview and .docx."""
    village = chain['village']
    lang = a.get('language', 'en')
    kind = a.get('kind', 'umuganda')
    partner = (a.get('partner') or '').strip()
    venue = (a.get('venue') or '').strip()
    gathering = (a.get('gathering_point') or '').strip()
    audience = (a.get('audience') or '').strip()
    start_time = (a.get('start_time') or '').strip()
    location_details = (a.get('location_details') or '').strip()
    reminder = (a.get('reminder') or '').strip()
    activities = (a.get('activities') or '').strip().rstrip('.')
    # Gathering points: the new list, or the single legacy pair
    points = [p for p in (a.get('meeting_points') or []) if isinstance(p, dict) and str(p.get('place', '')).strip()]
    if not points and gathering:
        points = [{'audience': audience, 'place': gathering}]
    bracket = f' ({location_details})' if location_details else ''

    def gathering_paragraph(default_who, will_gather, time_word):
        """One full sentence per gathering point: "<Who> will gather at <place> at <time>." """
        sentences = []
        for p in points:
            who = str(p.get('audience', '')).strip()
            s = f"{who} {will_gather}" if who else default_who
            s += f" {str(p['place']).strip()}"
            if str(p.get('time', '')).strip():
                s += f" {time_word} {str(p['time']).strip()}"
            sentences.append(s + '.')
        return ' '.join(sentences)
    note = (a.get('note') or '').strip()
    body = (a.get('body') or '').strip()
    event_date = a.get('event_date')

    out = []
    if lang == 'rw':
        title = (a.get('title') or '').strip() or 'ITANGAZO'
        out.append(('title', title.upper()))
        if body:
            out.extend(('body', p) for p in body.split('\n') if p.strip())
        elif kind == 'umuganda':
            text = f"Ubuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'bw')}"
            if partner:
                text += f' bufatanyije na {partner}'
            text += f' burabamenyesha ko hateganyijwe umuganda rusange kuwa {rw_date(event_date)}'
            if venue:
                text += f' ahitwa {venue}'
            text += '.'
            if start_time:
                text += f' Umuganda uzatangira saa {start_time}.'
            out.append(('body', text))
            # Where and when people gather, then what will be done (or that details follow on arrival)
            second = gathering_paragraph('Twese tuzahurira', 'bazahurira', 'saa')
            if activities:
                second += f" Ibikorwa bizakorwa ni: {activities}{bracket}."
            else:
                second += f" Andi makuru ku hantu n'ibikorwa azatangwa aho tuzahurira{bracket}."
            out.append(('body', second.strip()))
            if reminder:
                out.append(('body', reminder))
        else:
            text = f"Ubuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'bw')} burabamenyesha ko hateganyijwe " \
                   f"{(a.get('title') or 'inama').strip()} kuwa {rw_date(event_date)}"
            if venue:
                text += f' ahitwa {venue}'
            if start_time:
                text += f' guhera saa {start_time}'
            text += '.'
            if audience:
                text += f' Abatumiwe: {audience}.'
            out.append(('body', text))
        if note:
            out.append(('note', f'Icyitonderwa: {note}'))
        out.append(('sign_title', f"Umuyobozi {rw_of(rw_of(village, 'umudugudu w'), 'w')}"))
    else:
        title = (a.get('title') or '').strip() or 'ANNOUNCEMENT'
        out.append(('title', title.upper()))
        if body:
            out.extend(('body', p) for p in body.split('\n') if p.strip())
        elif kind == 'umuganda':
            text = f'The Office of {village} Village'
            if partner:
                text += f', in collaboration with {partner},'
            text += f' wishes to inform you of forthcoming community work (Umuganda) on {en_long_date(event_date)}'
            if venue:
                text += f' at {venue}'
            text += '.'
            if start_time:
                text += f' The community work (Umuganda) will start at {start_time}.'
            out.append(('body', text))
            # Where and when people gather, then what will be done (or that details follow on arrival)
            second = gathering_paragraph('We will all gather at', 'will gather at', 'at')
            if activities:
                second += f" The work will focus on {activities}{bracket}."
            else:
                second += f" Further details about the location and activities will be shared at the gathering point{bracket}."
            out.append(('body', second.strip()))
            if reminder:
                out.append(('body', reminder))
        else:
            text = f"The Office of {village} Village wishes to inform you of a forthcoming " \
                   f"{(a.get('title') or 'meeting').strip().lower()} on {en_long_date(event_date)}"
            if venue:
                text += f' at {venue}'
            if start_time:
                text += f', starting at {start_time}'
            text += '.'
            if audience:
                text += f' Invited: {audience}.'
            out.append(('body', text))
        if note:
            out.append(('note', f'Note: {note}'))
        out.append(('sign_title', 'Village Leader'))
    out.append(('sign_name', leader_name))
    return out


def render_announcement(ann, leader):
    chain = location_chain(ann.village or (ann.created_by.village if ann.created_by else None))
    data = {
        'language': ann.language, 'kind': ann.kind, 'title': ann.title, 'partner': ann.partner,
        'venue': ann.venue, 'gathering_point': ann.gathering_point, 'audience': ann.audience,
        'meeting_points': ann.meeting_points, 'location_details': ann.location_details, 'reminder': ann.reminder,
        'activities': ann.activities,
        'start_time': ann.start_time, 'note': ann.note, 'body': ann.body, 'event_date': ann.event_date,
    }
    info = leader_info(leader or ann.created_by)
    paragraphs = announcement_paragraphs(data, chain, info['name'])

    doc = _new_document()
    date_text = rw_date(ann.letter_date) if ann.language == 'rw' else en_date(ann.letter_date)
    _letterhead(doc, chain, date_text, ann.language)
    # Letterhead already ends with one blank line; five in total before the title, three after it
    _blank_lines(doc, 4)
    for style, text in paragraphs:
        if style == 'title':
            p = _title(doc, text, center=True)
            p.paragraph_format.space_after = Pt(0)
            _blank_lines(doc, 3)
        elif style == 'body':
            _para(doc, text)
        elif style == 'note':
            p = _para(doc, text)
            for run in p.runs:
                run.italic = True
        elif style == 'sign_title':
            doc.add_paragraph()
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = Pt(0)
        elif style == 'sign_name':
            p = doc.add_paragraph()
            p.add_run(text).bold = True
    if info['phone']:
        doc.add_paragraph(f"Tel: {info['phone']}")
    slug = (ann.title or ann.get_kind_display()).replace(' ', '_')[:40]
    return _save(doc), f'{slug}_{ann.letter_date}.docx'
